"""Create Arabic DOCX test files from existing book text."""
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_docx_from_txt(txt_path, output_path, title, author, max_paragraphs=100):
    """Create a DOCX file from Arabic text.

    TXT files from OCR/scans often have hard line wraps without blank-line
    paragraph separators. We treat each non-empty, non-page-number line as
    its own paragraph for DOCX testing purposes.
    """
    with open(txt_path, 'r', encoding='utf-8') as f:
        raw_text = f.read()

    # Each line becomes a paragraph (TXT books have hard wraps, no blank-line separators)
    lines = raw_text.strip().split('\n')
    paragraphs = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        # Skip page numbers (pure digits)
        if stripped.isdigit():
            continue
        paragraphs.append(stripped)

    print(f"  Source paragraphs: {len(paragraphs)}")

    # Create the DOCX
    doc = Document()

    # Set default style
    style = doc.styles['Normal']
    style.font.size = Pt(14)
    style.font.name = 'Arial'

    # Add title
    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Add author
    author_para = doc.add_paragraph(author)
    author_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author_para.runs[0].font.size = Pt(16)

    # Add page break
    doc.add_page_break()

    # Add paragraphs
    count = min(max_paragraphs, len(paragraphs))
    for i in range(count):
        p = doc.add_paragraph(paragraphs[i])
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.save(output_path)
    size = os.path.getsize(output_path)
    print(f"  Saved: {output_path}")
    print(f"  Size: {size:,} bytes")
    print(f"  Paragraphs: {count}")
    return size


def create_multi_format_docx(output_path):
    """Create a DOCX with various formatting for testing edge cases."""
    doc = Document()

    style = doc.styles['Normal']
    style.font.size = Pt(14)
    style.font.name = 'Arial'

    # Title
    h = doc.add_heading('نموذج اختبار التنسيقات المتعددة', level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Section 1: Regular paragraphs
    doc.add_heading('القسم الأول: فقرات عادية', level=1)
    texts = [
        'هذه فقرة عادية باللغة العربية. تحتوي على جملتين أو أكثر لاختبار تقسيم الفقرات.',
        'الفقرة الثانية تحتوي على نص أطول. نحتاج إلى التأكد من أن النظام يتعامل مع الفقرات المتعددة بشكل صحيح. هذا مهم جداً لعملية إنتاج الكتب الصوتية.',
    ]
    for t in texts:
        p = doc.add_paragraph(t)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Section 2: Dialogue
    doc.add_heading('القسم الثاني: حوار', level=1)
    dialogue_lines = [
        'دخل الرجل الغرفة ونظر حوله بحذر.',
        'قال: من أنت؟ وماذا تفعل هنا؟',
        'أجاب الشاب: أنا صديق ابنك. جئت أزوره.',
        'قال الرجل بصوت حاد: ابني ليس هنا. اذهب من حيث أتيت!',
        'لم يتحرك الشاب من مكانه. نظر إلى الرجل بثبات وقال: سأنتظره.',
    ]
    for line in dialogue_lines:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Section 3: Mixed content with bullet points
    doc.add_heading('القسم الثالث: قائمة', level=1)
    items = [
        'العنصر الأول في القائمة',
        'العنصر الثاني في القائمة',
        'العنصر الثالث في القائمة',
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    # Section 4: Poetry/verse (common in Arabic books)
    doc.add_heading('القسم الرابع: شعر', level=1)
    verses = [
        'قفا نبك من ذكرى حبيب ومنزل',
        'بسقط اللوى بين الدخول فحومل',
        'فتوضح فالمقراة لم يعف رسمها',
        'لما نسجتها من جنوب وشمأل',
    ]
    for v in verses:
        p = doc.add_paragraph(v)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(output_path)
    size = os.path.getsize(output_path)
    print(f"  Saved: {output_path}")
    print(f"  Size: {size:,} bytes")
    return size


if __name__ == '__main__':
    docx_dir = '/home/hamr/PycharmProjects/Sawt/data/books/docx/'
    os.makedirs(docx_dir, exist_ok=True)

    # 1. Create DOCX from existing TXT book
    print("Creating awlad-haretna.docx from TXT...")
    create_docx_from_txt(
        txt_path='/home/hamr/PycharmProjects/Sawt/data/books/txt/awalad-7aretna.txt',
        output_path=os.path.join(docx_dir, 'awlad-haretna.docx'),
        title='أولاد حارتنا',
        author='نجيب محفوظ',
        max_paragraphs=100,
    )

    # 2. Create DOCX from book2.txt
    print("\nCreating book2.docx from TXT...")
    create_docx_from_txt(
        txt_path='/home/hamr/PycharmProjects/Sawt/data/books/txt/book2.txt',
        output_path=os.path.join(docx_dir, 'book2.docx'),
        title='كتاب الاختبار الثاني',
        author='مؤلف مجهول',
        max_paragraphs=50,
    )

    # 3. Create a multi-format test DOCX
    print("\nCreating test-formats.docx (synthetic test file)...")
    create_multi_format_docx(os.path.join(docx_dir, 'test-formats.docx'))

    print("\nDone!")
